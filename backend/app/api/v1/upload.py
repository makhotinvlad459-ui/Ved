# backend/app/api/v1/upload.py
from fastapi import APIRouter, UploadFile, File, HTTPException, Depends, Form
from fastapi.responses import FileResponse
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select
import uuid
import os
from datetime import datetime
import shutil
import re
import tempfile

from app.database import get_db
from app.models import ProcessingSession, PendingModel
from app.celery.tasks import process_invoice

router = APIRouter(prefix="/upload", tags=["Upload"])

UPLOAD_DIR = "/app/uploads"
TEMPLATE_DIR = "/app/templates"
OUTPUT_DIR = "/app/output"
TEMP_DIR = "/app/temp"

os.makedirs(UPLOAD_DIR, exist_ok=True)
os.makedirs(TEMPLATE_DIR, exist_ok=True)
os.makedirs(OUTPUT_DIR, exist_ok=True)
os.makedirs(TEMP_DIR, exist_ok=True)


@router.post("/")
async def upload_files(
    invoice: UploadFile = File(...),
    manifest: UploadFile = File(...),
    template: UploadFile = File(...),
    spec_number: str = Form(default=""),
    db: AsyncSession = Depends(get_db)
):
    session_id = str(uuid.uuid4())
    
    invoice_path = os.path.join(UPLOAD_DIR, f"{session_id}_invoice.xlsx")
    manifest_path = os.path.join(UPLOAD_DIR, f"{session_id}_manifest.xlsx")
    template_path = os.path.join(TEMPLATE_DIR, f"{session_id}_template.xlsx")
    
    with open(invoice_path, "wb") as f:
        f.write(await invoice.read())
    with open(manifest_path, "wb") as f:
        f.write(await manifest.read())
    with open(template_path, "wb") as f:
        f.write(await template.read())
    
    session = ProcessingSession(
        session_id=session_id,
        status="pending",
        invoice_file=invoice_path,
        manifest_file=manifest_path,
        template_file=template_path,
        spec_number=spec_number,
        created_at=datetime.utcnow()
    )
    db.add(session)
    await db.commit()
    await db.refresh(session)
    
    process_invoice.delay(session_id)
    
    return {
        "session_id": session_id,
        "status": "pending",
        "message": "Файлы загружены, обработка запущена"
    }


@router.get("/status/{session_id}")
async def get_status(
    session_id: str,
    db: AsyncSession = Depends(get_db)
):
    stmt = select(ProcessingSession).where(ProcessingSession.session_id == session_id)
    result = await db.execute(stmt)
    session = result.scalar_one_or_none()
    
    if not session:
        raise HTTPException(status_code=404, detail="Сессия не найдена")
    
    pending_count = 0
    if session.status == "pending_approval":
        pending_stmt = select(PendingModel).where(
            PendingModel.session_id == session_id,
            PendingModel.status == "pending"
        )
        pending_result = await db.execute(pending_stmt)
        pending_count = len(pending_result.scalars().all())
    
    return {
        "session_id": session.session_id,
        "status": session.status,
        "spec_number": session.spec_number,
        "total_items": session.total_items,
        "processed_items": session.processed_items,
        "errors": session.errors,
        "result_file": session.result_file,
        "created_at": session.created_at,
        "invoice_date": session.invoice_date,
        "spec_date": session.spec_date,
        "pending_models_count": pending_count if session.status == "pending_approval" else 0,
        "message": f"Ожидается подтверждение {pending_count} новых моделей" if pending_count > 0 else None
    }


@router.get("/download/{session_id}")
async def download_result(
    session_id: str,
    db: AsyncSession = Depends(get_db)
):
    stmt = select(ProcessingSession).where(ProcessingSession.session_id == session_id)
    result = await db.execute(stmt)
    session = result.scalar_one_or_none()
    
    if not session:
        raise HTTPException(status_code=404, detail="Сессия не найдена")
    
    if session.status != "completed":
        raise HTTPException(status_code=400, detail="Обработка еще не завершена")
    
    result_path = session.result_file
    if not result_path or not os.path.exists(result_path):
        raise HTTPException(status_code=404, detail="Файл результата не найден")
    
    filename = os.path.basename(result_path)

    return FileResponse(
        path=result_path,
        filename=filename,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


@router.delete("/cleanup/{session_id}")
async def cleanup_session(
    session_id: str,
    db: AsyncSession = Depends(get_db)
):
    """Полная очистка всех файлов сессии"""
    stmt = select(ProcessingSession).where(ProcessingSession.session_id == session_id)
    result = await db.execute(stmt)
    session = result.scalar_one_or_none()
    
    if not session:
        raise HTTPException(status_code=404, detail="Сессия не найдена")
    
    files_to_delete = [
        session.invoice_file,
        session.manifest_file,
        session.template_file,
        session.result_file
    ]
    
    deleted = []
    errors = []
    
    for file_path in files_to_delete:
        if file_path and os.path.exists(file_path):
            try:
                os.remove(file_path)
                deleted.append(os.path.basename(file_path))
            except Exception as e:
                errors.append(f"{os.path.basename(file_path)}: {str(e)}")
    
    pending_stmt = select(PendingModel).where(PendingModel.session_id == session_id)
    pending_result = await db.execute(pending_stmt)
    pendings = pending_result.scalars().all()
    for p in pendings:
        p.status = "cleaned"
    
    session.status = "cleaned"
    await db.commit()
    
    return {
        "message": f"Удалено файлов: {len(deleted)}",
        "files": deleted,
        "errors": errors if errors else None,
        "session_id": session_id
    }


@router.get("/pending-count/{session_id}")
async def get_pending_count(
    session_id: str,
    db: AsyncSession = Depends(get_db)
):
    """Получить количество новых моделей, ожидающих подтверждения"""
    stmt = select(PendingModel).where(
        PendingModel.session_id == session_id,
        PendingModel.status == "pending"
    )
    result = await db.execute(stmt)
    pending_count = len(result.scalars().all())
    
    return {
        "session_id": session_id,
        "pending_models_count": pending_count,
        "has_pending": pending_count > 0
    }


# ============================================================
# НОВЫЙ ЭНДПОИНТ ДЛЯ ДЕКЛАРИРОВАНИЯ (ПОСТ + ОДИН ФАЙЛ)
# ============================================================
@router.post("/declaration-text")
async def get_declaration_text(
    invoice: UploadFile = File(...)
):
    from app.services.parser import parse_invoice

    with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
        tmp.write(await invoice.read())
        tmp_path = tmp.name

    try:
        invoice_data = parse_invoice(tmp_path)
        items = invoice_data.get("items", [])
        
        summary = {}
        
        for item in items:
            desc = item.get("description", "")
            model_num = item.get("model_number", "")
            qty = item.get("qty", 0)
            
            if not desc or qty <= 0:
                continue

            clean_name = desc.split(" (")[0].strip()
            clean_name = re.sub(r'^(Apple|Samsung|Google|Beats)\s+', '', clean_name, flags=re.IGNORECASE)
            
            # === 1. СПЕЦИАЛЬНЫЕ ПРАВИЛА ===
            
            # Чехлы, ремешки, адаптеры (аксессуары)
            if re.search(r'Case|Strap|Band|Adapter|Key Ring|Folio|Earbuds', clean_name, re.IGNORECASE):
                if "Band for Apple Watch" in clean_name:
                    clean_name = "Band for Apple Watch"
                elif "Crossbody Strap" in clean_name:
                    clean_name = "Crossbody Strap"
                elif "TechWoven Case" in clean_name:
                    clean_name = "TechWoven Case"
                elif "Smart Folio" in clean_name:
                    clean_name = "Smart Folio"
                elif "Key Ring" in clean_name:
                    clean_name = "FineWoven Key Ring"
                elif "Power Adapter" in clean_name:
                    m = re.search(r'\d+W\s+USB-C\s+Power Adapter', clean_name, re.IGNORECASE)
                    clean_name = m.group(0) if m else "Power Adapter"
                elif "AirPods" in clean_name:
                    clean_name = "AirPods with Wireless Charging Case" if "Wireless" in clean_name else "AirPods"
                
                clean_name = re.sub(r'\b\d+(?:Gb|GB|Tb|TB|mm)\b', '', clean_name, flags=re.IGNORECASE)
                clean_name = re.sub(r'\bwith\b.*$', '', clean_name, flags=re.IGNORECASE)
                clean_name = re.sub(r'^[\s\-\/,]+|[\s\-\/,]+$', '', clean_name).strip()

            # Часы (жёсткая обрезка по слову Aluminium)
            elif re.search(r'Watch (SE|Series)', clean_name, re.IGNORECASE):
                # Обрезаем всё, что идет после "Aluminium"
                clean_name = clean_name.split(" Aluminium")[0].strip()
                # Убираем лишние пробелы
                clean_name = re.sub(r'\s+', ' ', clean_name)

            # Fitbit
            elif re.search(r'Fitbit', clean_name, re.IGNORECASE):
                clean_name = "Google Fitbit Air"

            # iPad Pro
            elif re.search(r'iPad Pro \d+', clean_name, re.IGNORECASE):
                m = re.search(r'iPad Pro \d+["\u201d\u2033]?', clean_name, re.IGNORECASE)
                if m:
                    clean_name = m.group(0).replace('"', '').strip()
                else:
                    clean_name = "iPad Pro"
                    
            # MacBook Air / Pro
            elif re.search(r'MacBook (Air|Pro)', clean_name, re.IGNORECASE):
                m = re.search(r'MacBook (?:Air|Pro)\s+\d+["\u201d\u2033]?', clean_name, re.IGNORECASE)
                if m:
                    clean_name = m.group(0).replace('"', '').strip()
                    
            # Mac Mini
            elif "Mac Mini" in clean_name:
                clean_name = "Mac Mini"
                
            # iPhone
            elif "iPhone" in clean_name:
                if "Duo" in clean_name:
                    clean_name = "iPhone Duo"
                else:
                    m = re.search(r'iPhone\s+\d+(?:\s+(?:Pro\s+Max|Pro|Plus|Air|Mini|Max))?', clean_name, re.IGNORECASE)
                    if m:
                        clean_name = m.group(0)
                    else:
                        clean_name = "iPhone"

            # === 2. ОБЩАЯ ОЧИСТКА ===
            else:
                clean_name = re.sub(r'\b\d+(?:Gb|GB|Tb|TB|Гб|Тб|mm)\b', '', clean_name, flags=re.IGNORECASE)
                colors = ["Black", "White", "Silver", "Blue", "Green", "Yellow", "Pink", "Purple", 
                          "Gold", "Grey", "Gray", "Midnight", "Starlight", "Space", "Burgundy", 
                          "Red", "Orange", "Glacier", "Star", "Crisp", "Sand", "Taupe", "Olive"]
                for color in colors:
                    clean_name = re.sub(rf'\b{re.escape(color)}\b', '', clean_name, flags=re.IGNORECASE)
                clean_name = re.sub(r'\bwith\b.*$', '', clean_name, flags=re.IGNORECASE)
                clean_name = re.sub(r'\b\d+["\u201d\u2033]\b', '', clean_name, flags=re.IGNORECASE)
                clean_name = re.sub(r'\s+', ' ', clean_name).strip()
                clean_name = re.sub(r'^[\s\-\/,]+|[\s\-\/,]+$', '', clean_name).strip()

            if not clean_name or len(clean_name) < 2:
                clean_name = desc.split(" (")[0].strip()

            # === 3. ГРУППИРОВКА ===
            key = clean_name
            if key not in summary:
                summary[key] = {}
                
            if model_num:
                model_clean = " ".join(sorted(model_num.split()))
                summary[key][model_clean] = summary[key].get(model_clean, 0) + qty
            else:
                summary[key]["—"] = summary[key].get("—", 0) + qty

        lines = []
        for model_name, models_dict in summary.items():
            if len(models_dict) == 1 and "—" in models_dict:
                lines.append(f"{model_name} - {list(models_dict.values())[0]} шт.;")
            else:
                models_str = ", ".join([f"модель {m} - {q} шт." for m, q in sorted(models_dict.items()) if m != "—"])
                if "—" in models_dict:
                     models_str += f", без модели - {models_dict['—']} шт."
                lines.append(f"{model_name}, {models_str};")
            
        text = "\n".join(lines)
        return {"text": text}
        
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)

@router.post("/far-east-text")
async def get_far_east_text(
    file: UploadFile = File(...)
):
    import openpyxl
    import tempfile
    import re

    with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
        tmp.write(await file.read())
        tmp_path = tmp.name

    try:
        wb = openpyxl.load_workbook(tmp_path, data_only=True)
        ws = wb.active

        # 1. Заголовки
        header_row = None
        col_idx = {}
        for r in range(1, 15):
            row_vals = [str(ws.cell(r, c).value).lower().strip() if ws.cell(r, c).value else "" for c in range(1, 15)]
            if "наименование" in row_vals and "шт" in row_vals:
                header_row = r
                for c in range(1, 15):
                    val = str(ws.cell(r, c).value).lower().strip()
                    if "наименование" in val: col_idx["name"] = c
                    elif "шт" in val: col_idx["qty"] = c
                    elif "модель" in val: col_idx["model"] = c
                    elif "бренд" in val or "фирма" in val: col_idx["brand"] = c
                break

        if not header_row:
            raise Exception("Не найдены заголовки таблицы.")

        items = []
        for r in range(header_row + 1, ws.max_row + 1):
            name_val = ws.cell(r, col_idx.get("name")).value
            qty_val = ws.cell(r, col_idx.get("qty")).value
            model_val = ws.cell(r, col_idx.get("model")).value if "model" in col_idx else None
            brand_val = ws.cell(r, col_idx.get("brand")).value if "brand" in col_idx else None

            if not name_val or not qty_val: continue
            try: qty = int(qty_val)
            except: continue
            if qty <= 0: continue

            items.append({
                "name": str(name_val).strip(),
                "qty": qty,
                "model": str(model_val).strip() if model_val else "",
                "brand": str(brand_val).strip() if brand_val else "",
            })

        # Словарь замены кириллицы на латиницу (для моделей и названий)
        CYR_TO_LAT = {
            'А': 'A', 'В': 'B', 'С': 'C', 'Е': 'E', 'Н': 'H', 'К': 'K', 'М': 'M',
            'О': 'O', 'Р': 'P', 'Т': 'T', 'Х': 'X', 'У': 'Y', 'а': 'a', 'е': 'e',
            'о': 'o', 'р': 'p', 'с': 'c', 'у': 'y', 'х': 'x'
        }
        def fix_cyr(s):
            return "".join(CYR_TO_LAT.get(c, c) for c in s)

        grouped = {}
        
        for item in items:
            name = item["name"]
            model = fix_cyr(item["model"])
            qty = item["qty"]
            brand = item["brand"]

            # Очистка названия
            clean_name = fix_cyr(name)
            
            # Убираем IMEI и мусор
            clean_name = re.sub(r'(imei\d*|imai\d*|серия|модель,)\s*[\d\/]+', '', clean_name, flags=re.IGNORECASE)
            clean_name = re.sub(r'\b\d{10,}\b', '', clean_name)
            clean_name = re.sub(r'\b\d+[/\s+]*\d*\s*(gb|тб|tb|гб)\b', '', clean_name, flags=re.IGNORECASE)
            clean_name = re.sub(r'^(смартфон|планшетный компьютер|планшет|наушники|смарт-часы|умные часы|часы|ноутбук|игровая приставка|игровая консоль|выпрямитель|фен для волос|фен|стайлер|музыкальная колонка)\s+', '', clean_name, flags=re.IGNORECASE)
            clean_name = re.sub(r'^(смартфон|планшетный компьютер|планшет|наушники|смарт-часы|умные часы|часы|ноутбук|игровая приставка|игровая консоль|выпрямитель|фен для волос|фен|стайлер|музыкальная колонка)\s+', '', clean_name, flags=re.IGNORECASE)

            # Очистка цветов
            colors = ["Black", "White", "Silver", "Blue", "Green", "Yellow", "Pink", "Purple", "Gold", "Grey", "Gray", "Midnight", "Starlight", "Space", "Burgundy", "Red", "Orange", "Glacier", "Star", "Crisp", "Sand", "Taupe", "Olive", "Teal", "Ultramarine", "Mist", "Deep", "Cosmic", "Rose", "Navy", "Titanium", "Graphite", "Strawberry", "Bronze", "Blush", "Ceramic", "Jasper", "Plum", "Kanzan", "Sky", "Shadow", "Sage", "Orane", "Amber", "Silk", "Copper", "Dark", "Light"]
            for c in colors:
                clean_name = re.sub(rf'\b{re.escape(c)}\b', '', clean_name, flags=re.IGNORECASE)

            # Убираем бренды
            clean_name = re.sub(r'^(Apple|Samsung|Sony|Dyson|Xiaomi|Redmi|JBL|Yandex|Яндекс|Nintendo|Google|Fitbit)\s+', '', clean_name, flags=re.IGNORECASE)
            clean_name = re.sub(r'\s+', ' ', clean_name).strip()
            clean_name = re.sub(r'[,\-\.]+$', '', clean_name).strip()

            # --- СПЕЦИАЛЬНЫЕ ПРАВИЛА (Формируем чистое имя) ---
            lower_name = clean_name.lower()
            
            if "iphone" in lower_name or "айфон" in lower_name:
                m = re.search(r'iphone\s*\d*(?:\s+(?:pro\s+max|pro|plus|air|mini|max))?', clean_name, re.IGNORECASE)
                if m: clean_name = re.sub(r'\s+', ' ', m.group(0)).title().replace("Iphone", "iPhone")
                else: clean_name = "iPhone"
            
            elif "ipad" in lower_name or "айпад" in lower_name:
                m = re.search(r'ipad\s+(?:pro|air|mini)?\s*\d*', clean_name, re.IGNORECASE)
                if m: clean_name = m.group(0).title().replace("Ipad", "iPad").replace('"', '').strip()
            
            elif "macbook" in lower_name:
                m = re.search(r'macbook\s+(?:air|pro)\s*\d*', clean_name, re.IGNORECASE)
                if m: clean_name = m.group(0).title().replace("Macbook", "MacBook").replace('"', '').strip()
            
            elif "airpods" in lower_name or "air pods" in lower_name:
                m = re.search(r'air\s*pods\s*(?:pro|max)?\s*\d*', clean_name, re.IGNORECASE)
                if m: clean_name = m.group(0).title().replace("Airpods", "AirPods").replace("Air Pods", "AirPods")
            
            elif "watch" in lower_name or "воч" in lower_name:
                m = re.search(r'watch\s*(?:se|series)?\s*\d+', clean_name, re.IGNORECASE)
                if m: clean_name = m.group(0).title().replace("Watch", "Watch")
            
            elif "galaxy" in lower_name:
                m = re.search(r'galaxy\s*[a-z]?\d*(?:\s+(?:ultra|pro|fe|plus))?', clean_name, re.IGNORECASE)
                if m: clean_name = m.group(0).title().replace("Galaxy", "Galaxy")
                elif "buds" in lower_name: clean_name = "Galaxy Buds"
            
            elif "xiaomi" in lower_name or "redmi" in lower_name:
                # Оставляем Pad или 13 Pro
                m = re.search(r'(?:redmi\s+|xiaomi\s+)?(?:pad|13|12s)\s*[a-z0-9]*', clean_name, re.IGNORECASE)
                if m: clean_name = m.group(0).title().replace("Xiaomi", "Xiaomi").replace("Redmi", "Redmi")
                elif "фен" in lower_name: clean_name = "Xiaomi Фен"
                elif "ирригатор" in lower_name: clean_name = "Xiaomi Ирригатор"
            
            elif "dyson" in lower_name or brand.upper() == "DYSON":
                if "supersonic" in lower_name or "hd16" in model or "hd17" in model: clean_name = "Dyson Supersonic"
                elif "airwrap" in lower_name or "airw" in lower_name or "hs05" in model or "hs08" in model or "hs09" in model: clean_name = "Dyson Airwrap"
                elif "airstrait" in lower_name or "airwait" in lower_name or "ht01" in model: clean_name = "Dyson Airstrait"
                elif "detect" in lower_name or "sv23" in model: clean_name = "Dyson Detect"
                else: clean_name = "Dyson"
            
            elif "playstation" in lower_name or "ps5" in lower_name:
                if "джойстик" in lower_name or "dual" in lower_name: clean_name = "Sony DualSense"
                elif "pro" in lower_name: clean_name = "Sony PS5 Pro"
                else: clean_name = "Sony PS5"
            
            elif "яндекс" in lower_name or "yandex" in lower_name or "станция" in lower_name:
                if "макс" in lower_name: clean_name = "Yandex Станция Макс"
                elif "мини" in lower_name and "pro" in lower_name: clean_name = "Yandex Станция Мини Pro"
                elif "мини" in lower_name: clean_name = "Yandex Станция Мини"
                elif "миди" in lower_name: clean_name = "Yandex Станция Миди"
                elif "дуо" in lower_name: clean_name = "Yandex Станция Дуо"
                elif "стрит" in lower_name: clean_name = "Yandex Станция Стрит"
                else: clean_name = "Yandex Станция"
            
            elif "nintendo" in lower_name:
                clean_name = "Nintendo Switch OLED"
            
            # Если осталось пустое
            if not clean_name or len(clean_name) < 2:
                clean_name = name.split("(")[0].split(",")[0].strip()

            # --- ОПРЕДЕЛЕНИЕ БРЕНДА ---
            b_up = brand.upper() if brand else ""
            if "APPLE" in b_up: brand_clean = "APPLE"
            elif "SAMSUNG" in b_up: brand_clean = "SAMSUNG"
            elif "SONY" in b_up: brand_clean = "SONY"
            elif "DYSON" in b_up: brand_clean = "DYSON"
            elif "XIAOMI" in b_up: brand_clean = "XIAOMI"
            elif "JBL" in b_up or "HARMAN" in b_up: brand_clean = "JBL"
            elif "YANDEX" in b_up or "ЯНДЕКС" in b_up: brand_clean = "YANDEX"
            elif "NINTENDO" in b_up: brand_clean = "NINTENDO"
            else:
                if "iphone" in lower_name or "ipad" in lower_name or "macbook" in lower_name: brand_clean = "APPLE"
                elif "samsung" in lower_name or "galaxy" in lower_name: brand_clean = "SAMSUNG"
                elif "dyson" in lower_name: brand_clean = "DYSON"
                elif "xiaomi" in lower_name or "redmi" in lower_name: brand_clean = "XIAOMI"
                elif "sony" in lower_name or "playstation" in lower_name: brand_clean = "SONY"
                elif "jbl" in lower_name: brand_clean = "JBL"
                elif "яндекс" in lower_name or "yandex" in lower_name or "станция" in lower_name: brand_clean = "YANDEX"
                elif "nintendo" in lower_name: brand_clean = "NINTENDO"
                else: brand_clean = "ПРОЧЕЕ"

            model_key = "—" if (not model or model.lower() == "none") else " ".join(sorted(model.split()))

            if brand_clean not in grouped: grouped[brand_clean] = {}
            if clean_name not in grouped[brand_clean]: grouped[brand_clean][clean_name] = {}
            grouped[brand_clean][clean_name][model_key] = grouped[brand_clean][clean_name].get(model_key, 0) + qty

        # Формируем текст
        lines = []
        for brand, models_dict in sorted(grouped.items()):
            lines.append(f"=== {brand} ===")
            for model_name, model_nums in sorted(models_dict.items()):
                if len(model_nums) == 1 and "—" in model_nums:
                    lines.append(f"{model_name} - {list(model_nums.values())[0]} шт.;")
                else:
                    parts = []
                    for m, q in sorted(model_nums.items()):
                        if m == "—": parts.append(f"без модели - {q} шт.")
                        else: parts.append(f"модель {m} - {q} шт.")
                    lines.append(f"{model_name}, {', '.join(parts)};")
            lines.append("")

        return {"text": "\n".join(lines).strip()}

    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)