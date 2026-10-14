# backend/app/services/for_alice_generator.py
"""
Генератор xlsx для сервиса For Alice.
"""
import os
from typing import Dict, Any, List
from decimal import Decimal, ROUND_HALF_UP

import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter


OUTPUT_DIR = "/app/output"


def _round2(v):
    if v is None:
        return None
    return float(Decimal(str(v)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def generate_for_alice_xlsx(
    result: Dict[str, Any],
    session_id: str,
) -> str:
    """
    Генерирует xlsx с 14 колонками.
    """
    os.makedirs(OUTPUT_DIR, exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Расчёт"

    # === Стили ===
    header_font = Font(bold=True, size=10)
    header_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # === Заголовки ===
    headers = [
        "Код модели",
        "P/N",
        "Наименование",
        "Кол-во",
        "Вес",
        "Вес +15%",
        "Общий вес",
        "Цена (AED)",
        "Стоимость партии",
        "Логистика/ед",
        "Цена + логистика",
        "Стоимость с логистикой",
        "+ЧЗ (AED)",
        "Итог",
    ]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(1, col)
        cell.value = h
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = border

    # === Данные ===
    row = 2
    for r in result["rows"]:
        values = [
            r.get("model_number") or "",
            r.get("part_number") or "",
            r.get("name") or "",
            r.get("qty") or 0,
            _round2(r.get("weight")),
            _round2(r.get("g")),
            _round2(r.get("h")),
            _round2(r.get("price")),
            _round2((r.get("price") or 0) * (r.get("qty") or 0)),
            _round2(r.get("l")),
            _round2(r.get("m")),
            _round2(r.get("n")),
            _round2(r.get("o")),
            _round2(r.get("p")),
        ]
        for col, v in enumerate(values, 1):
            cell = ws.cell(row, col)
            cell.value = v
            cell.border = border
            if col == 3:
                cell.alignment = left_align
            else:
                cell.alignment = cell_align
            if col in (5, 6, 7, 8, 9, 10, 11, 12, 13, 14):
                cell.number_format = "#,##0.00"
            if col == 4:
                cell.number_format = "#,##0"
        row += 1

    # === Итоги по таблице ===
    row += 1
    ws.cell(row, 1).value = "Итого вес:"
    ws.cell(row, 1).font = Font(bold=True)
    ws.cell(row, 2).value = _round2(result["totals"]["total_weight"])
    ws.cell(row, 2).number_format = "#,##0.00"
    row += 1

    ws.cell(row, 1).value = "Итого сумма:"
    ws.cell(row, 1).font = Font(bold=True)
    ws.cell(row, 2).value = _round2(result["totals"]["total_sum"])
    ws.cell(row, 2).number_format = "#,##0.00"
    row += 2

    # === Параметры расчёта (ввод) ===
    ws.cell(row, 1).value = "Параметры расчёта:"
    ws.cell(row, 1).font = Font(bold=True)
    row += 1

    ws.cell(row, 1).value = "Ставка логистики (USD/кг)"
    ws.cell(row, 2).value = result["totals"]["logistics_rate_usd"]
    ws.cell(row, 2).number_format = "#,##0.00"
    row += 1

    ws.cell(row, 1).value = "Курс USD → AED"
    ws.cell(row, 2).value = result["totals"]["usd_to_aed"]
    ws.cell(row, 2).number_format = "#,##0.000"
    row += 1

    ws.cell(row, 1).value = "Расходы на документы (AED)"
    ws.cell(row, 2).value = result["totals"]["documents_cost_aed"]
    ws.cell(row, 2).number_format = "#,##0.00"
    row += 2

    # === Расчёт логистики ===
    ws.cell(row, 1).value = "Расчёт логистики:"
    ws.cell(row, 1).font = Font(bold=True)
    row += 1

    ws.cell(row, 1).value = "Логистика всего (USD)"
    ws.cell(row, 2).value = _round2(result["totals"]["logistics_total_usd"])
    ws.cell(row, 2).number_format = "#,##0.00"
    row += 1

    ws.cell(row, 1).value = "Логистика всего (AED)"
    ws.cell(row, 2).value = _round2(result["totals"]["logistics_total_aed"])
    ws.cell(row, 2).number_format = "#,##0.00"
    row += 1

    ws.cell(row, 1).value = "Итого логистика (AED)"
    ws.cell(row, 2).value = _round2(result["totals"]["total_logistics_aed"])
    ws.cell(row, 2).number_format = "#,##0.00"
    row += 1

    ws.cell(row, 1).value = "Ставка AED/кг (с документами)"
    ws.cell(row, 2).value = _round2(result["totals"]["rate_aed_per_kg"])
    ws.cell(row, 2).number_format = "#,##0.00"
    row += 1

    ws.cell(row, 1).value = "Ставка USD/кг (с документами)"
    ws.cell(row, 2).value = _round2(result["totals"]["rate_usd_per_kg"])
    ws.cell(row, 2).number_format = "#,##0.00"

    # === Ширины колонок ===
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 60
    ws.column_dimensions["D"].width = 8
    for col in range(5, 15):
        ws.column_dimensions[get_column_letter(col)].width = 14

    output_path = os.path.join(OUTPUT_DIR, f"for_alice_{session_id[:8]}.xlsx")
    wb.save(output_path)

    return output_path