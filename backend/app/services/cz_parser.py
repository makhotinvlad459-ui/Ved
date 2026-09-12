import openpyxl
from typing import Dict, List, Any
import re


def parse_cz_file(file_path: str) -> Dict[str, List[str]]:
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb.active
    
    groups = {}
    col_code = None
    col_gtin = None
    
    for row in range(1, min(ws.max_row, 20)):
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row, col)
            if cell.value and isinstance(cell.value, str):
                cell_lower = cell.value.lower()
                if \"код\" in cell_lower or \"code\" in cell_lower:
                    col_code = col
                elif \"gtin\" in cell_lower:
                    col_gtin = col
        if col_code and col_gtin:
            break
    
    if not col_code or not col_gtin:
        raise Exception(\"Не удалось найти колонки 'Код' и 'GTIN'\")
    
    for row in range(2, ws.max_row + 1):
        code = ws.cell(row, col_code).value
        gtin = ws.cell(row, col_gtin).value
        
        if not code or not gtin:
            continue
        
        code = str(code).strip()
        gtin = str(gtin).strip()
        gtin = re.sub(r'[^0-9]', '', gtin)
        
        if not gtin:
            continue
        
        if gtin not in groups:
            groups[gtin] = []
        groups[gtin].append(code)
    
    return groups