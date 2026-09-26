# backend/app/services/chestny_znak_generator.py
import openpyxl
from openpyxl.styles import Alignment
import re
from typing import Dict, List, Optional


def parse_spec_info(file_path: str):
    """
    Извлекает из спецификации:
      - номер спецификации (например, "109") или None
      - номер строки заголовков таблицы
      - номер строки с Total AED
    """
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active

    spec_number: Optional[str] = None
    header_row: Optional[int] = None
    total_row: Optional[int] = None

    # Ищем "№ X от Y" в шапке (строки 1..25)
    for row in range(1, min(26, ws.max_row + 1)):
        a_val = ws.cell(row, 1).value
        c_val = ws.cell(row, 3).value
        if a_val and "СПЕЦИФИКАЦИЯ" in str(a_val) and c_val:
            m = re.search(r'№\s*([^\s]+)', str(c_val))
            if m:
                spec_number = m.group(1).strip()

        # Заголовок таблицы: A = "№" и B содержит "Описание"/"Description"
        if a_val and str(a_val).strip() == "№":
            b_val = ws.cell(row, 2).value
            if b_val and ("Описание" in str(b_val) or "Description" in str(b_val)):
                header_row = row

    # Ищем строку Total AED
    for row in range(1, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            v = ws.cell(row, col).value
            if v and "Total AED" in str(v):
                total_row = row
                break
        if total_row:
            break

    return spec_number, header_row, total_row


def build_invoice_index(invoice_data: dict) -> Dict[str, List[dict]]:
    """{part_number: [item, item, ...]} с сохранением порядка."""
    idx: Dict[str, List[dict]] = {}
    for item in invoice_data.get("items", []):
        pn = item.get("part_number")
        if pn:
            idx.setdefault(pn, []).append(item)
    return idx


def generate_cz_specification(
    spec_template_path: str,
    output_path: str,
    cz_codes: Dict[str, List[str]],
    invoice_data: dict,
) -> str:
    spec_number, header_row, total_row = parse_spec_info(spec_template_path)
    if not header_row:
        raise Exception("Не найден заголовок таблицы в спецификации")

    wb = openpyxl.load_workbook(spec_template_path)
    ws = wb.active

    # 1. Переименование заголовка колонки F
    new_header = f"ЧЗ - код №{spec_number}" if spec_number else "ЧЗ - код"
    for row in range(1, header_row + 1):
        cell = ws.cell(row, 6)
        if cell.value and isinstance(cell.value, str):
            v = cell.value
            if "Серийный" in v or "Serial" in v:
                cell.value = new_header
                cell.alignment = Alignment(
                    wrap_text=True, horizontal='center', vertical='center'
                )

    # 2. Индекс инвойса
    invoice_idx = build_invoice_index(invoice_data)

    # 3. Глобальные счётчики + защита от повторного использования
    used: Dict[str, int] = {gtin: 0 for gtin in cz_codes}
    consumed: set = set()   # все уже выданные коды (на всякий случай)

    # 4. Обход строк данных
    end_row = total_row if total_row else ws.max_row + 1
    row = header_row + 1

    while row < end_row:
        part_number_cell = ws.cell(row, 5).value
        if not part_number_cell:
            row += 1
            continue

        part_number = str(part_number_cell).strip()

        qty_val = ws.cell(row, 8).value
        try:
            qty = int(qty_val) if qty_val is not None else 0
        except (ValueError, TypeError):
            qty = 0

        gtin: Optional[str] = None
        inv_items = invoice_idx.get(part_number, [])
        if inv_items:
            for it in inv_items:
                upc_digits = re.sub(r'\D', '', str(it.get("upc") or ""))
                if upc_digits:
                    gtin = upc_digits.zfill(14)
                    break

        # Собираем коды для строки
        codes_for_row: List[str] = []
        if gtin and gtin in cz_codes and qty > 0:
            pool = cz_codes[gtin]
            start = used[gtin]
            # Набираем qty кодов, пропуская уже выданные
            for code in pool[start:]:
                if len(codes_for_row) >= qty:
                    break
                if code in consumed:
                    continue
                codes_for_row.append(code)
                consumed.add(code)
                used[gtin] += 1

        # Добиваем N/A, если не хватило
        if qty > 0 and len(codes_for_row) < qty:
            codes_for_row.extend(["N/A"] * (qty - len(codes_for_row)))

        if not codes_for_row:
            codes_for_row = ["N/A"]

        cell = ws.cell(row, 6)
        cell.value = "\n".join(codes_for_row)
        cell.alignment = Alignment(
            wrap_text=True, vertical='top', horizontal='left'
        )

        row += 1

    wb.save(output_path)
    return output_path