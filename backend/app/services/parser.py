# backend/app/services/parser.py
import openpyxl
from datetime import datetime
from typing import List, Dict, Any, Optional


def parse_invoice(file_path: str) -> Dict[str, Any]:
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb.active
    
    # === Дата и номер ===
    invoice_date = None
    invoice_number = None
    
    for row in ws.iter_rows(max_row=30):
        for cell in row:
            if cell.value and isinstance(cell.value, str):
                val = cell.value.strip()
                
                if "Reference" in val:
                    ref_cell = ws.cell(row=cell.row, column=14)
                    if ref_cell.value:
                        invoice_number = str(ref_cell.value).strip()
                        print(f"   📝 Найден номер инвойса: {invoice_number}")
                
                if val == "Date":
                    date_cell = ws.cell(row=cell.row, column=14)
                    if date_cell.value:
                        if isinstance(date_cell.value, datetime):
                            invoice_date = date_cell.value
                        elif isinstance(date_cell.value, str):
                            try:
                                invoice_date = datetime.strptime(date_cell.value, "%d.%m.%Y")
                            except:
                                try:
                                    invoice_date = datetime.strptime(date_cell.value, "%Y-%m-%d")
                                except:
                                    pass
                        print(f"   📝 Найдена дата инвойса: {invoice_date}")
                
                if invoice_date and invoice_number:
                    break
        if invoice_date and invoice_number:
            break
    
    # === Находим ВСЕ заголовки таблиц ===
    header_rows = []
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        has_model = False
        has_pn = False
        for cell in row:
            if cell.value and isinstance(cell.value, str):
                val = cell.value.strip()
                if val == "Model":
                    has_model = True
                elif "P/N" in val:
                    has_pn = True
        if has_model and has_pn:
            header_rows.append(row[0].row)
    
    print(f"   📋 Найдено блоков: {len(header_rows)} (строки: {header_rows})")
    
    items = []
    
    for idx, header_row in enumerate(header_rows):
        # ЖЁСТКО ЗАШИВАЕМ КОЛОНКИ (инвойс K625 всегда одинаковый)
        # Model: C (3), P/N: E (5), Description: H (8), Qty: V (22), COO: AC (29), UPC: AD (30), Price: AE (31), Total: AI (35)
        header_cells = {
            "model": 3,
            "part": 5,
            "desc": 8,
            "qty": 22,
            "coo": 29,
            "upc": 30,
            "price": 31,
            "total": 35
        }
        
        print(f"   📋 Блок {idx+1} (строка {header_row}): {header_cells}")
        
        if idx + 1 < len(header_rows):
            end_row = header_rows[idx + 1] - 1
        else:
            end_row = ws.max_row
        
        for row in ws.iter_rows(min_row=header_row + 1, max_row=end_row):
            # P/N в колонке E (5) -> index 4
            part_idx = 4
            part_cell = row[part_idx] if part_idx < len(row) else None
            
            if not part_cell or not part_cell.value:
                continue
            
            part_number = str(part_cell.value).strip()
            
            if part_number.lower() in ["total", "итого", "p/n", "part"]:
                continue
            
            # Model в колонке C (3) -> index 2
            model_idx = 2
            model_cell = row[model_idx] if model_idx < len(row) else None
            model_number = str(model_cell.value).strip() if model_cell and model_cell.value else ""
            
            # Description в колонке H (8) -> index 7
            desc_idx = 7
            desc_cell = row[desc_idx] if desc_idx < len(row) else None
            description = str(desc_cell.value).strip() if desc_cell and desc_cell.value else ""
            
            # Qty в колонке V (22) -> index 21
            qty_idx = 21
            qty_cell = row[qty_idx] if qty_idx < len(row) else None
            qty = 0
            if qty_cell and qty_cell.value:
                try:
                    qty = int(qty_cell.value)
                except (ValueError, TypeError):
                    qty = 0
            
            # COO в колонке AC (29) -> index 28
            coo_idx = 28
            coo_cell = row[coo_idx] if coo_idx < len(row) else None
            coo = str(coo_cell.value).strip() if coo_cell and coo_cell.value else ""
            
            # UPC в колонке AD (30) -> index 29
            upc_idx = 29
            upc_cell = row[upc_idx] if upc_idx < len(row) else None
            upc = str(upc_cell.value).strip() if upc_cell and upc_cell.value else ""
            
            # Price в колонке AE (31) -> index 30
            price_idx = 30
            price_cell = row[price_idx] if price_idx < len(row) else None
            price = 0.0
            if price_cell and price_cell.value:
                try:
                    price = float(price_cell.value)
                except (ValueError, TypeError):
                    price = 0.0
            
            # Total в колонке AI (35) -> index 34
            total_idx = 34
            total_cell = row[total_idx] if total_idx < len(row) else None
            total = 0.0
            if total_cell and total_cell.value:
                try:
                    total = float(total_cell.value)
                except (ValueError, TypeError):
                    total = 0.0
            
            # Пропускаем только строки с заголовками (где part_number = "P/N" или "Model")
            if part_number in ["P/N", "Model", "N/A", "Description"] and qty == 0:
                continue

            items.append({
                "model_number": model_number,
                "part_number": part_number,
                "description": description,
                "qty": qty,
                "coo": coo,
                "upc": upc,
                "price": price,
                "total": total
            })
    
    print(f"   ✅ Всего позиций: {len(items)}")
    
    return {
        "invoice_date": invoice_date,
        "invoice_number": invoice_number,
        "items": items
    }


def parse_manifest(file_path: str) -> Dict[str, List[str]]:
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb.active

    serials_by_part: Dict[str, List[str]] = {}
    imei_by_part: Dict[str, List[str]] = {}

    header_row = None
    for row in ws.iter_rows(min_row=1, max_row=10):
        for cell in row:
            if cell.value and isinstance(cell.value, str):
                val = cell.value.strip()
                if "Part Number" in val or val == "Part":
                    header_row = cell.row
                    break
        if header_row:
            break

    if header_row:
        part_col = None
        serial_col = None
        imei_col = None

        for col_idx, cell in enumerate(ws[header_row], 1):
            if cell.value is None:
                continue
            val = str(cell.value).strip()
            v_low = val.lower()

            if v_low == "part number" or v_low == "part":
                part_col = col_idx
            elif "serial" in v_low and serial_col is None:
                serial_col = col_idx
            elif "imei" in v_low and imei_col is None:
                imei_col = col_idx

        print(f"   📋 Манифест: part_col={part_col}, serial_col={serial_col}, "
              f"imei_col={imei_col} (только IMEI 1)")

        for row in ws.iter_rows(min_row=header_row + 1):
            if part_col is None or part_col - 1 >= len(row):
                continue
            part_cell = row[part_col - 1]
            if not part_cell or not part_cell.value:
                continue
            part_number = str(part_cell.value).strip()
            if not part_number:
                continue

            if part_number not in serials_by_part:
                serials_by_part[part_number] = []
                imei_by_part[part_number] = []

            if serial_col is not None and serial_col - 1 < len(row):
                serial_val = row[serial_col - 1].value
                if serial_val:
                    serials_by_part[part_number].append(str(serial_val).strip())

            if imei_col is not None and imei_col - 1 < len(row):
                imei_val = row[imei_col - 1].value
                if imei_val:
                    imei_by_part[part_number].append(str(imei_val).strip())

    else:
        for row in ws.iter_rows(min_row=2):
            part_number = row[3].value
            if not part_number:
                break
            part_key = str(part_number).strip()

            if part_key not in serials_by_part:
                serials_by_part[part_key] = []
                imei_by_part[part_key] = []

            if row[4].value:
                serials_by_part[part_key].append(str(row[4].value).strip())
            if row[5].value:
                imei_by_part[part_key].append(str(row[5].value).strip())

    print(f"   ✅ Серийников: {sum(len(v) for v in serials_by_part.values())}, "
          f"IMEI (только 1): {sum(len(v) for v in imei_by_part.values())}")

    return {
        "serials": serials_by_part,
        "imei": imei_by_part
    }