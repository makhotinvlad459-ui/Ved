# backend/app/services/packing_list_parser.py
import openpyxl
from typing import List, Dict, Any, Optional, Tuple
from datetime import datetime
import re


def parse_packing_list(file_path: str) -> Dict[str, Any]:
    """
    Парсинг файла Packing List
    """
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb.active
    
    header_row = None
    headers = {}
    
    for row in range(1, min(ws.max_row, 20)):
        for col in range(1, ws.max_column + 1):
            cell_value = ws.cell(row, col).value
            if cell_value and isinstance(cell_value, str):
                if "Pallet no." in cell_value or "Pallet" in cell_value:
                    header_row = row
                    break
        if header_row:
            break
    
    if not header_row:
        for row in range(1, min(ws.max_row, 20)):
            for col in range(1, ws.max_column + 1):
                cell_value = ws.cell(row, col).value
                if cell_value and isinstance(cell_value, str):
                    if "Description" in cell_value:
                        header_row = row
                        break
            if header_row:
                break
    
    if not header_row:
        raise Exception("Не удалось найти заголовки таблицы")
    
    for col in range(1, ws.max_column + 1):
        value = ws.cell(header_row, col).value
        if value and isinstance(value, str):
            value_clean = value.strip().lower()
            if "pallet" in value_clean:
                headers["pallet_no"] = col
            elif "box" in value_clean:
                headers["box_no"] = col
            elif "description" in value_clean or "items" in value_clean:
                headers["description"] = col
            elif "qty" in value_clean or "quantity" in value_clean:
                headers["qty"] = col
            elif "weight per 1" in value_clean or "weight per" in value_clean:
                headers["weight_per_1"] = col
            elif "net" in value_clean:
                headers["net"] = col
            elif "gross" in value_clean:
                headers["gross"] = col
            elif "weight (kgs)" in value_clean or "weight" in value_clean:
                headers["pallet_weight"] = col
            elif "dimensions" in value_clean:
                headers["dimensions"] = col
    
    items = []
    current_pallet = None
    current_box = None
    
    for row in range(header_row + 1, ws.max_row + 1):
        row_empty = True
        for col in range(1, ws.max_column + 1):
            if ws.cell(row, col).value:
                row_empty = False
                break
        
        if row_empty:
            continue
        
        pallet_no = ws.cell(row, headers.get("pallet_no", 1)).value if "pallet_no" in headers else None
        box_no = ws.cell(row, headers.get("box_no", 2)).value if "box_no" in headers else None
        description = ws.cell(row, headers.get("description", 3)).value if "description" in headers else None
        qty = ws.cell(row, headers.get("qty", 4)).value if "qty" in headers else None
        weight_per_1 = ws.cell(row, headers.get("weight_per_1", 5)).value if "weight_per_1" in headers else None
        net = ws.cell(row, headers.get("net", 6)).value if "net" in headers else None
        gross = ws.cell(row, headers.get("gross", 7)).value if "gross" in headers else None
        pallet_weight = ws.cell(row, headers.get("pallet_weight", 8)).value if "pallet_weight" in headers else None
        dimensions = ws.cell(row, headers.get("dimensions", 9)).value if "dimensions" in headers else None
        
        if pallet_no:
            current_pallet = pallet_no
        if box_no:
            current_box = box_no
        
        if description and not description.isdigit():
            part_number = str(description).strip()
            part_number = re.sub(r'\)$', '', part_number)
            part_number = part_number.strip()
            
            items.append({
                "part_number": part_number,
                "qty": int(qty) if qty else 0,
                "weight_per_1": float(weight_per_1) if weight_per_1 else None,
                "net": float(net) if net else None,
                "gross": float(gross) if gross else None,
                "pallet_weight": float(pallet_weight) if pallet_weight else None,
                "dimensions": str(dimensions) if dimensions else None,
                "pallet_no": current_pallet,
                "box_no": current_box
            })
    
    return {"items": items}


def get_product_by_part_number(db, part_number: str) -> Tuple[Optional[Any], Optional[str]]:
    """
    Ищет продукт по part_number (с нормализацией).
    Возвращает (product, model_number).
    """
    from app.models import Product

    if not part_number:
        return None, None

    # Нормализация: убираем пробелы, приводим к верхнему регистру
    normalized = part_number.strip().upper()

    product = db.query(Product).filter(Product.part_number == normalized).first()

    if product:
        return product, product.model_number

    # Fallback: попробовать как есть (на случай разных регистров в БД)
    product = db.query(Product).filter(Product.part_number == part_number).first()

    if product:
        return product, product.model_number

    return None, None


def redistribute_gross(items: List[Dict], pallet_weight: float) -> List[Dict]:
    """
    Выравнивание GROSS между позициями с учётом 19% ограничения
    """
    if not items or pallet_weight <= 0:
        return items
    
    total_net = sum(item.get("net", 0) or 0 for item in items)
    if total_net == 0:
        return items
    
    for item in items:
        net = item.get("net", 0) or 0
        item["gross"] = (net / total_net) * pallet_weight
    
    max_iterations = 100
    for _ in range(max_iterations):
        diffs = []
        for item in items:
            net = item.get("net", 0) or 0
            gross = item.get("gross", 0) or 0
            if net > 0:
                diff = (gross - net) / net * 100
                diffs.append((item, diff, net, gross))
            else:
                diffs.append((item, 0, net, gross))
        
        max_diff_item, max_diff, max_net, max_gross = max(diffs, key=lambda x: x[1])
        min_diff_item, min_diff, min_net, min_gross = min(diffs, key=lambda x: x[1])
        
        if max_diff - min_diff < 0.5:
            break
        
        if max_diff <= 19 and min_diff >= 0:
            break
        
        transfer = 0.1
        
        if min_gross - transfer >= min_net:
            min_diff_item["gross"] = min_gross - transfer
        else:
            min_diff_item["gross"] = min_net
        
        max_diff_item["gross"] = max_gross + transfer
        
        new_max_gross = max_diff_item.get("gross", 0)
        if max_net > 0:
            new_max_diff = (new_max_gross - max_net) / max_net * 100
            if new_max_diff > 19:
                max_diff_item["gross"] = max_gross
                min_diff_item["gross"] = min_gross
                break
    
    for item in items:
        net = item.get("net", 0) or 0
        gross = item.get("gross", 0) or 0
        if net > 0:
            diff = (gross - net) / net * 100
            if diff > 19:
                new_net = gross / 1.19
                item["net"] = new_net
    
    for item in items:
        net = item.get("net", 0) or 0
        qty = item.get("qty", 0) or 0
        gross = item.get("gross", 0) or 0
        
        if qty > 0:
            item["weight_per_item"] = net / qty
            item["total_net"] = net
            item["total_gross"] = gross
            if net > 0:
                item["difference_percent"] = (gross - net) / net * 100
            else:
                item["difference_percent"] = 0
    
    return items