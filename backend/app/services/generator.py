# backend/app/services/generator.py
import openpyxl
from openpyxl.styles import Alignment
from datetime import datetime, timedelta
from typing import List, Dict, Any
import re


# Регулярки для «мягких» пробелов
WS = r'[\s\xa0]'          # whitespace или неразрывный пробел
NUM_DATE = r'[\d\.]+'     # 25.08.2026


def generate_specification(
    template_path: str,
    data: List[Dict[str, Any]],
    spec_number: str,
    invoice_date: datetime,
    invoice_number: str,
    output_path: str
) -> str:
    """
    Генерация спецификации по шаблону
    """
    wb = openpyxl.load_workbook(template_path)
    ws = wb.active

    # ============================================
    # 1. Заменяем номера и даты в шапке
    # ============================================
    spec_date = invoice_date - timedelta(days=1) if invoice_date else datetime.now()
    date_str = spec_date.strftime("%d.%m.%Y")
    invoice_date_str = invoice_date.strftime("%d.%m.%Y") if invoice_date else date_str
    new_spec_number = spec_number if spec_number and spec_number.strip() else "105"

    print(f"   📝 Номер спецификации: {new_spec_number}, дата: {date_str}")
    print(f"   📝 Номер инвойса: {invoice_number}, дата инвойса: {invoice_date_str}")

    for row in range(1, 25):
        # ---------- РУССКАЯ ЧАСТЬ ----------
        ru_spec_cell = ws.cell(row, 1)
        ru_number_cell = ws.cell(row, 3)

        if ru_spec_cell.value and "СПЕЦИФИКАЦИЯ" in str(ru_spec_cell.value):
            if ru_number_cell.value and "№" in str(ru_number_cell.value):
                val = str(ru_number_cell.value).strip()
                # Пример: "№ 105 от 25.08.2026"
                match = re.search(
                    rf'№{WS}*([^\s\xa0]+){WS}+от{WS}+({NUM_DATE})',
                    val
                )
                if match:
                    old_number = match.group(1)
                    new_val = re.sub(
                        rf'№{WS}*{re.escape(old_number)}',
                        f'№ {new_spec_number}',
                        val,
                        count=1
                    )
                    new_val = re.sub(
                        rf'от{WS}+{NUM_DATE}',
                        f'от {date_str}',
                        new_val,
                        count=1
                    )
                    ru_number_cell.value = new_val
                    print(f"   📝 [RU] Номер спецификации: {old_number} → {new_spec_number}, дата → {date_str}")

        # ---------- АНГЛИЙСКАЯ ЧАСТЬ ----------
        en_spec_cell = ws.cell(row, 7)
        en_number_cell = ws.cell(row, 8)

        if en_spec_cell.value and "SPECIFICATION" in str(en_spec_cell.value):
            if en_number_cell.value and "№" in str(en_number_cell.value):
                val = str(en_number_cell.value).strip()
                # Пример: "№ 105 dt. 25.08.2026"
                match = re.search(
                    rf'№{WS}*([^\s\xa0]+){WS}+dt\.{WS}*({NUM_DATE})',
                    val
                )
                if match:
                    old_number = match.group(1)
                    new_val = re.sub(
                        rf'№{WS}*{re.escape(old_number)}',
                        f'№ {new_spec_number}',
                        val,
                        count=1
                    )
                    new_val = re.sub(
                        rf'dt\.{WS}*{NUM_DATE}',
                        f'dt. {date_str}',
                        new_val,
                        count=1
                    )
                    en_number_cell.value = new_val
                    print(f"   📝 [EN] Номер спецификации: {old_number} → {new_spec_number}, дата → {date_str}")

        # ---------- НОМЕР ИНВОЙСА (русская часть) ----------
        ru_invoice_cell = ws.cell(row, 1)
        ru_invoice_number_cell = ws.cell(row, 3)

        if ru_invoice_cell.value and "Коммерческому инвойсу" in str(ru_invoice_cell.value):
            if ru_invoice_number_cell.value and "№" in str(ru_invoice_number_cell.value):
                val = str(ru_invoice_number_cell.value).strip()
                match = re.search(
                    rf'№{WS}*([^\s\xa0]+){WS}+от{WS}+({NUM_DATE})',
                    val
                )
                if match:
                    old_invoice = match.group(1)
                    new_val = re.sub(
                        rf'№{WS}*{re.escape(old_invoice)}',
                        f'№ {invoice_number}',
                        val,
                        count=1
                    )
                    new_val = re.sub(
                        rf'от{WS}+{NUM_DATE}',
                        f'от {invoice_date_str}',
                        new_val,
                        count=1
                    )
                    ru_invoice_number_cell.value = new_val
                    print(f"   📝 [RU] Номер инвойса: {old_invoice} → {invoice_number}, дата → {invoice_date_str}")

        # ---------- НОМЕР ИНВОЙСА (английская часть) ----------
        en_invoice_cell = ws.cell(row, 7)
        en_invoice_number_cell = ws.cell(row, 8)

        if en_invoice_cell.value and "Commercial Invoice" in str(en_invoice_cell.value):
            if en_invoice_number_cell.value and "№" in str(en_invoice_number_cell.value):
                val = str(en_invoice_number_cell.value).strip()
                match = re.search(
                    rf'№{WS}*([^\s\xa0]+){WS}+dt\.{WS}*({NUM_DATE})',
                    val
                )
                if match:
                    old_invoice = match.group(1)
                    new_val = re.sub(
                        rf'№{WS}*{re.escape(old_invoice)}',
                        f'№ {invoice_number}',
                        val,
                        count=1
                    )
                    new_val = re.sub(
                        rf'dt\.{WS}*{NUM_DATE}',
                        f'dt. {invoice_date_str}',
                        new_val,
                        count=1
                    )
                    en_invoice_number_cell.value = new_val
                    print(f"   📝 [EN] Номер инвойса: {old_invoice} → {invoice_number}, дата → {invoice_date_str}")

    # ============================================
    # 2. Находим строку с заголовками таблицы
    # ============================================
    header_row = None
    for row in range(1, min(ws.max_row, 50)):
        cell_a = ws.cell(row, 1)
        cell_b = ws.cell(row, 2)
        if cell_a.value == "№" or (cell_a.value and str(cell_a.value).strip() == "№"):
            if cell_b.value and ("Описание" in str(cell_b.value) or "Description" in str(cell_b.value)):
                header_row = row
                break

    if not header_row:
        header_row = 15
        print(f"   ⚠️ Заголовок не найден, используем строку {header_row}")

    print(f"   📍 Заголовок таблицы: строка {header_row}")

    # ============================================
    # 3. Находим строку с итогами
    # ============================================
    total_row = None
    for row in range(header_row + 1, min(ws.max_row, 100)):
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row, col)
            if cell.value and "Total AED" in str(cell.value):
                total_row = row
                break
        if total_row:
            break

    if not total_row:
        for row in range(1, ws.max_row + 1):
            for col in range(1, ws.max_column + 1):
                cell = ws.cell(row, col)
                if cell.value and "Total AED" in str(cell.value):
                    total_row = row
                    break
            if total_row:
                break

    if total_row:
        print(f"   📍 Строка с итогами: {total_row}")

    # ============================================
    # 4. Очищаем старые данные
    # ============================================
    start_row = header_row + 1
    end_row = total_row - 1 if total_row else ws.max_row

    if start_row <= end_row:
        for row in range(start_row, end_row + 1):
            for col in range(1, 11):
                cell = ws.cell(row, col)
                is_merged = False
                for merged_range in ws.merged_cells.ranges:
                    if cell.coordinate in merged_range:
                        is_merged = True
                        break
                if not is_merged:
                    cell.value = None

    # ============================================
    # 5. Вставляем новые данные
    # ============================================
    current_row = start_row

    for idx, item in enumerate(data):
        if total_row and current_row >= total_row:
            ws.insert_rows(current_row)
            total_row += 1

        serials_value = item.get('serials', '')
        if serials_value == "Не заполняется":
            serials_value = ""

        for col, value in {
            1: idx + 1,
            2: item.get('name', ''),
            3: item.get('model', ''),
            4: item.get('coo', ''),
            5: item.get('article', ''),
            7: '',
            8: item.get('qty', 0),
            9: item.get('price', 0),
            10: item.get('total', 0)
        }.items():
            cell = ws.cell(current_row, col)
            is_merged = False
            for merged_range in ws.merged_cells.ranges:
                if cell.coordinate in merged_range:
                    is_merged = True
                    break
            if not is_merged:
                cell.value = value
                if col == 2:
                    cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='left')
                else:
                    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

        # Запись серийников
        f_cell = ws.cell(current_row, 6)
        target_cell = f_cell
        for merged_range in ws.merged_cells.ranges:
            if f_cell.coordinate in merged_range:
                target_cell = ws.cell(merged_range.min_row, merged_range.min_col)
                break

        target_cell.value = serials_value
        target_cell.alignment = Alignment(wrap_text=True, vertical='top', horizontal='left')

        if serials_value:
            print(f"   📝 Строка {current_row}: записано {len(serials_value.split(chr(10)))} серийников")

        ws.row_dimensions[current_row].height = 30
        current_row += 1

    # ============================================
    # 6. Обновляем формулы сумм
    # ============================================
    if total_row:
        h_cell = ws.cell(total_row, 8)
        if h_cell.value and "SUM" in str(h_cell.value):
            h_cell.value = f"=SUM(H{start_row}:H{current_row - 1})"

        j_cell = ws.cell(total_row, 10)
        if j_cell.value and "SUM" in str(j_cell.value):
            j_cell.value = f"=SUM(J{start_row}:J{current_row - 1})"

    wb.save(output_path)
    print(f"   ✅ Спецификация сохранена: {output_path}")
    print(f"   📊 Записано {len(data)} строк")

    return output_path