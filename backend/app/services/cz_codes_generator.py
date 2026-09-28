# backend/app/services/cz_codes_generator.py
import openpyxl
from openpyxl.styles import Font, Alignment
import re
import os
import zipfile
from typing import Optional, List, Dict, Tuple


def detect_group(description: str) -> Optional[str]:
    """
    Определяет группу по описанию.
      MacBook  → "MacBook"
      iPad     → "iPad"
      iPhone   → "Smartphone"
      Остальное → None (пропускаем)
    """
    if not description:
        return None
    d = description.lower()
    if "macbook" in d:
        return "MacBook"
    if "ipad" in d:
        return "iPad"
    if "iphone" in d:
        return "Smartphone"
    return None


def normalize_country(country: str) -> str:
    """
    Нормализует страну для имени файла:
      "Vietnam"  → "Vietnam"
      "China"    → "China"
      " UAE "    → "UAE"
      None       → "Unknown"
    """
    if not country:
        return "Unknown"
    s = str(country).strip()
    if not s:
        return "Unknown"
    # Пробелы → подчёркивания, убираем лишние символы
    s = re.sub(r'\s+', '_', s)
    return s


def parse_cz_spec(file_path: str) -> Tuple[Optional[str], List[dict]]:
    """
    Читает файл "ЧЗ - код №XXX.xlsx".
    Возвращает (spec_number, rows) где rows:
      [{"description": "...", "country": "China", "codes": ["010...", ...]}, ...]

    Из колонки F (ЧЗ-код) берёт все строки, разделённые \n,
    выкидывает пустые и "N/A".
    """
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb.active

    spec_number: Optional[str] = None

    # Ищем номер спецификации в шапке ("№ 109 от ...")
    for row in range(1, 12):
        a_val = ws.cell(row, 1).value
        c_val = ws.cell(row, 3).value
        if a_val and c_val and "СПЕЦИФИКАЦИЯ" in str(a_val):
            m = re.search(r'№\s*([^\s]+)', str(c_val))
            if m:
                spec_number = m.group(1).strip()
                break

    # Заголовок таблицы (A = "№", B содержит "Описание")
    header_row = None
    for row in range(1, min(ws.max_row, 50) + 1):
        a_val = ws.cell(row, 1).value
        b_val = ws.cell(row, 2).value
        if a_val and str(a_val).strip() == "№":
            if b_val and ("Описание" in str(b_val) or "Description" in str(b_val)):
                header_row = row
                break
    if not header_row:
        raise Exception("Не найден заголовок таблицы в файле ЧЗ")

    # Строка Total
    total_row = None
    for row in range(header_row + 1, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            v = ws.cell(row, col).value
            if v and "Total AED" in str(v):
                total_row = row
                break
        if total_row:
            break

    end_row = total_row if total_row else ws.max_row + 1

    rows: List[dict] = []
    for row in range(header_row + 1, end_row):
        desc_val = ws.cell(row, 2).value
        country_val = ws.cell(row, 4).value
        codes_val = ws.cell(row, 6).value

        if not desc_val:
            continue

        codes: List[str] = []
        if codes_val:
            for line in str(codes_val).split("\n"):
                code = line.strip()
                if not code:
                    continue
                if code.upper() == "N/A":
                    continue
                codes.append(code)

        rows.append({
            "description": str(desc_val).strip(),
            "country": str(country_val).strip() if country_val else "Unknown",
            "codes": codes,
        })

    return spec_number, rows


def group_codes(rows: List[dict]) -> Dict[Tuple[str, str], List[str]]:
    """
    Группирует коды по (группа, страна).
    Возвращает {(group, country): [code1, code2, ...]}
    """
    grouped: Dict[Tuple[str, str], List[str]] = {}
    for row in rows:
        group = detect_group(row["description"])
        if not group:
            continue
        country = normalize_country(row["country"])
        key = (group, country)
        if key not in grouped:
            grouped[key] = []
        grouped[key].extend(row["codes"])
    return grouped


def write_codes_xlsx(codes: List[str], output_path: str) -> str:
    """
    Пишет простой список кодов в столбец A, одна строка = один код.
    Без шапки, без нумерации — для загрузки в ЧЗ.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Коды"

    for idx, code in enumerate(codes, start=1):
        cell = ws.cell(idx, 1)
        cell.value = code
        cell.font = Font(name="Calibri", size=11)
        cell.alignment = Alignment(horizontal='left', vertical='center')

    ws.column_dimensions['A'].width = 40

    wb.save(output_path)
    return output_path


def generate_cz_codes_zip(
    cz_spec_path: str,
    output_dir: str,
    session_id: str,
) -> Tuple[str, Optional[str], Dict[str, int]]:
    """
    Читает ЧЗ-спецификацию, группирует коды, создаёт отдельные xlsx
    и упаковывает их в ZIP.

    Возвращает (zip_path, spec_number, {filename: count})
    """
    spec_number, rows = parse_cz_spec(cz_spec_path)
    grouped = group_codes(rows)

    if not grouped:
        raise Exception("Не найдено ни одной позиции MacBook/iPad/iPhone с кодами ЧЗ")

    # Рабочая папка для временных xlsx
    work_dir = os.path.join(output_dir, f"{session_id}_cz_parts")
    os.makedirs(work_dir, exist_ok=True)

    files_info: Dict[str, int] = {}
    generated_paths: List[str] = []

    for (group, country), codes in grouped.items():
        if not codes:
            continue
        filename = f"{group}_{country}.xlsx"
        file_path = os.path.join(work_dir, filename)
        write_codes_xlsx(codes, file_path)
        generated_paths.append(file_path)
        files_info[filename] = len(codes)

    if not generated_paths:
        raise Exception("Не найдено ни одного кода для выгрузки")

    # ZIP
    spec_suffix = spec_number if spec_number else session_id[:8]
    zip_name = f"ЧЗ_ввод_№{spec_suffix}.zip"
    zip_path = os.path.join(output_dir, zip_name)

    with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as zf:
        for file_path in generated_paths:
            # Внутри архива — только имя файла, без путей
            zf.write(file_path, arcname=os.path.basename(file_path))

    # Чистим временные xlsx
    for file_path in generated_paths:
        try:
            os.remove(file_path)
        except Exception:
            pass
    try:
        os.rmdir(work_dir)
    except Exception:
        pass

    return zip_path, spec_number, files_info