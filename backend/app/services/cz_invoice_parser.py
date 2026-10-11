# backend/app/services/cz_invoice_parser.py
"""
Парсер инвойса для сервиса «Карточки ЧЗ».
Извлекает позиции, определяет категорию, парсит атрибуты.
"""
import openpyxl
import re
from typing import List, Dict, Any, Optional


# ============================================================
# ОПРЕДЕЛЕНИЕ КАТЕГОРИИ
# ============================================================

def detect_category(description: str) -> Optional[str]:
    """Возвращает 'Смартфоны' / 'Планшеты' / 'Ноутбуки' / None."""
    if not description:
        return None
    d = description.lower()

    if "iphone" in d:
        return "Смартфоны"
    if "samsung" in d and "galaxy" in d:
        return "Смартфоны"
    if "ipad" in d:
        return "Планшеты"
    if "macbook" in d:
        return "Ноутбуки"

    return None


# ============================================================
# БРЕНД
# ============================================================

def detect_brand(description: str) -> Optional[str]:
    if not description:
        return None
    d = description.lower()
    if d.startswith("apple") or " apple " in d or "(apple" in d:
        return "APPLE"
    if "samsung" in d:
        return "Samsung"
    if "beats" in d:
        return "Beats"
    return None


# ============================================================
# АРТИКУЛ
# ============================================================

def extract_part_number_from_description(description: str) -> Optional[str]:
    if not description:
        return None
    matches = re.findall(r'\(([^()]+)\)', str(description))
    if matches:
        candidate = matches[-1].strip()
        if re.fullmatch(r'[A-Za-z0-9\-/. ]+', candidate):
            return candidate.upper()
    return None


# ============================================================
# АТРИБУТЫ
# ============================================================

def extract_screen_size(description: str) -> Optional[float]:
    if not description:
        return None

    m = re.search(r'(\d+(?:[.,]\d+)?)\s*["\u201d\u2033]', description)
    if m:
        try:
            return float(m.group(1).replace(",", "."))
        except ValueError:
            pass

    m = re.search(r'\b(?:ipad|macbook)\s+\w*\s*(\d{1,2})\b', description, re.IGNORECASE)
    if m:
        try:
            val = float(m.group(1))
            if 6 <= val <= 20:
                return val
        except ValueError:
            pass

    return None


def extract_ram_gb(description: str) -> Optional[int]:
    if not description:
        return None

    m = re.search(r'/\s*(\d+)\s*(?:Gb|GB|Гб)\s*/', description)
    if m:
        try:
            return int(m.group(1))
        except ValueError:
            pass

    m = re.search(r'GPU\s*/\s*(\d+)\s*(?:Gb|GB|Гб)', description, re.IGNORECASE)
    if m:
        try:
            return int(m.group(1))
        except ValueError:
            pass

    m = re.search(r'\b(\d+)\s*/\s*(\d+)\s*(?:Gb|GB|Гб|Tb|TB|Тб)', description)
    if m:
        try:
            return int(m.group(1))
        except ValueError:
            pass

    return None


def extract_rom_gb(description: str) -> Optional[int]:
    if not description:
        return None

    m = re.search(r'\b(\d+(?:[.,]\d+)?)\s*(?:Tb|TB|Тб)\b', description)
    if m:
        try:
            val = float(m.group(1).replace(",", "."))
            return int(val * 1024)
        except ValueError:
            pass

    m = re.search(r'\b(\d+)\s*(?:Gb|GB|Гб)\s*SSD', description, re.IGNORECASE)
    if m:
        try:
            return int(m.group(1))
        except ValueError:
            pass

    m = re.search(r'\b\d+\s*/\s*(\d+)\s*(?:Gb|GB|Гб)', description)
    if m:
        try:
            return int(m.group(1))
        except ValueError:
            pass

    m = re.search(r'\b(\d+)\s*(?:Gb|GB|Гб)\b', description)
    if m:
        try:
            val = int(m.group(1))
            if val >= 32:
                return val
        except ValueError:
            pass

    return None


def extract_cpu(description: str) -> tuple:
    if not description:
        return None, None

    m = re.search(r'(Snapdragon\s+\d+(?:\s+Elite(?:\s+Gen\s+\d+)?)?)', description, re.IGNORECASE)
    if m:
        return "Qualcomm", m.group(1).strip()

    m = re.search(r'(Exynos\s+\d+)', description, re.IGNORECASE)
    if m:
        return "Samsung", m.group(1).strip()

    m = re.search(r'\b(M\d+)\s+(Pro|Max|Ultra)\b', description)
    if m:
        return "Apple", f"{m.group(1)} {m.group(2)}"
    m = re.search(r'\b(M\d+)\b', description)
    if m:
        return "Apple", m.group(1)

    m = re.search(r'\b(A\d+)\s+(Pro|Max|Ultra)\b', description)
    if m:
        return "Apple", f"{m.group(1)} {m.group(2)}"
    m = re.search(r'\b(A\d+)\b', description)
    if m:
        return "Apple", m.group(1)

    return None, None


# ============================================================
# UPC → GTIN
# ============================================================

def upc_to_gtin(upc: Any) -> Optional[str]:
    if upc is None:
        return None
    s = re.sub(r'\D', '', str(upc))
    if not s:
        return None
    if len(s) == 12:
        return "0" + s
    if len(s) == 13:
        return s
    if len(s) == 14:
        return s[-13:]
    if len(s) < 12:
        return s.zfill(13)
    return s[-13:]


# ============================================================
# ПАРСИНГ ИНВОЙСА
# ============================================================

def parse_invoice_for_cz(file_path: str) -> List[Dict[str, Any]]:
    """
    Читает инвойс, возвращает список позиций со всеми атрибутами.
    Пропускает позиции без UPC и без категории ЧЗ.
    """
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb.active

    header_rows = []
    for row_idx in range(1, ws.max_row + 1):
        has_hash = False
        has_pn = False
        for col in range(1, ws.max_column + 1):
            v = ws.cell(row_idx, col).value
            if v and isinstance(v, str):
                s = v.strip()
                if s == "#":
                    has_hash = True
                elif s == "P/N":
                    has_pn = True
        if has_hash and has_pn:
            header_rows.append(row_idx)

    if not header_rows:
        raise Exception("Не найдена таблица в инвойсе (# / P/N)")

    print(f"   📋 Найдено блоков таблиц: {len(header_rows)}")

    items: List[Dict[str, Any]] = []

    for block_idx, header_row in enumerate(header_rows):
        cols: Dict[str, int] = {}
        for col in range(1, ws.max_column + 1):
            v = ws.cell(header_row, col).value
            if not v or not isinstance(v, str):
                continue
            s = v.strip()
            if s == "#":
                cols["num"] = col
            elif s == "Model":
                cols["model"] = col
            elif s == "P/N":
                cols["part"] = col
            elif s == "Description":
                cols["desc"] = col
            elif s == "Qty":
                cols["qty"] = col
            elif s == "COO":
                cols["coo"] = col
            elif "UPC" in s:
                cols["upc"] = col

        if block_idx + 1 < len(header_rows):
            end_row = header_rows[block_idx + 1] - 1
        else:
            end_row = ws.max_row

        for row_idx in range(header_row + 1, end_row + 1):
            def get_val(key):
                col = cols.get(key)
                if not col:
                    return None
                return ws.cell(row_idx, col).value

            num_val = get_val("num")
            part_val = get_val("part")
            desc_val = get_val("desc")

            if not desc_val:
                continue

            desc_str = str(desc_val).strip()
            if desc_str.upper().startswith("TOTAL"):
                continue

            part_number = None
            if part_val:
                part_number = str(part_val).strip().upper()
            if not part_number:
                part_number = extract_part_number_from_description(desc_str)
            if not part_number:
                continue

            category = detect_category(desc_str)
            if not category:
                continue

            upc_val = get_val("upc")
            gtin = upc_to_gtin(upc_val)
            if not gtin:
                continue

            brand = detect_brand(desc_str)

            qty_val = get_val("qty")
            try:
                qty = int(qty_val) if qty_val else 0
            except (ValueError, TypeError):
                qty = 0

            coo_val = get_val("coo")
            coo = str(coo_val).strip() if coo_val else None

            model_val = get_val("model")
            model_number = str(model_val).strip() if model_val else None

            cpu_vendor, cpu_model = extract_cpu(desc_str)

            items.append({
                "num": num_val,
                "part_number": part_number,
                "model_number": model_number,
                "description": desc_str,
                "brand": brand,
                "category": category,
                "qty": qty,
                "coo": coo,
                "upc": str(upc_val).strip() if upc_val else None,
                "gtin": gtin,
                "screen_size": extract_screen_size(desc_str),
                "ram_gb": extract_ram_gb(desc_str),
                "rom_gb": extract_rom_gb(desc_str),
                "cpu_vendor": cpu_vendor,
                "cpu_model": cpu_model,
            })

    print(f"   ✅ Распознано позиций: {len(items)}")
    return items