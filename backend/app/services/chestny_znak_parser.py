# backend/app/services/chestny_znak_parser.py
import openpyxl
import re
from typing import Dict, List


def parse_cz_codes(file_path: str) -> Dict[str, List[str]]:
    """
    Парсит файл кодов Честного Знака.
    Возвращает {gtin_14: [code1, code2, ...]}.

    Поддерживает два формата:
      1) Sheet0 — табличный: колонки "Код" и "GTIN".
      2) Лист1 — pivot: строка с GTIN, под ней — коды.
    """
    wb = openpyxl.load_workbook(file_path, data_only=True)

    # === Вариант 1: Sheet0 с колонками "Код" / "GTIN" ===
    if "Sheet0" in wb.sheetnames:
        ws = wb["Sheet0"]
        code_col = gtin_col = header_row = None

        for row in range(1, min(ws.max_row, 20) + 1):
            for col in range(1, ws.max_column + 1):
                v = ws.cell(row, col).value
                if v and isinstance(v, str):
                    val = v.strip().lower()
                    if val == "код":
                        code_col = col
                    elif val == "gtin":
                        gtin_col = col
            if code_col and gtin_col:
                header_row = row
                break

        if header_row and code_col and gtin_col:
            result: Dict[str, List[str]] = {}
            for row in range(header_row + 1, ws.max_row + 1):
                code_raw = ws.cell(row, code_col).value
                gtin_raw = ws.cell(row, gtin_col).value
                if not code_raw or not gtin_raw:
                    continue
                code = str(code_raw).strip()
                if not code or code.lower() in ("общий итог", "итого", "total"):
                    continue
                gtin_digits = re.sub(r'\D', '', str(gtin_raw))
                if not gtin_digits:
                    continue
                gtin = gtin_digits.zfill(14)
                result.setdefault(gtin, []).append(code)

            if result:
                return result

    # === Вариант 2: pivot-лист ===
    result: Dict[str, List[str]] = {}
    for sheet_name in wb.sheetnames:
        if sheet_name == "Sheet0":
            continue
        ws = wb[sheet_name]
        current_gtin = None
        for row in range(1, ws.max_row + 1):
            v = ws.cell(row, 1).value
            if not v:
                continue
            s = str(v).strip()
            if not s:
                continue
            if re.fullmatch(r'\d{13,14}', s):
                current_gtin = s.zfill(14)
                result.setdefault(current_gtin, [])
            elif current_gtin and s.startswith("01"):
                result[current_gtin].append(s)

    return result