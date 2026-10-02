# backend/app/services/commercial_offer_generator.py
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from datetime import datetime
import re
from typing import Optional, Tuple, List


# ============================================================
# ЖЁСТКО ЗАШИТЫЕ РЕКВИЗИТЫ (константы)
# ============================================================
SUPPLIER = {
    "name_ru": "KVN Group FZCO",
    "name_en": "KVN Group FZCO",
    "address_ru": "6WA 131, First Floor, 6 West A, Dubai Airport, UAE",
    "address_en": "6WA 131, First Floor, 6 West A, Dubai Airport, UAE",
    "phone_ru": "tel: +971 50 258 2030",
    "phone_en": "tel: +971 50 258 2030",
}

BUYER = {
    "name_ru": "ООО «КомБытТех»",
    "name_en": "LLC KomBytTech",
    "address_ru_1": "Россия, г.Москва",
    "address_ru_2": "ш Алтуфьевское, дом 79А, строение 3, помещение 6/3/2",
    "address_en_1": "Russia, Moscow",
    "address_en_2": "79А Altufevskoe Highway, building 3, office 6/3/2",
    "phone_ru": "Тел.: +7 926 639 46 99",
    "phone_en": "tel: +7 926 639 46 99",
}

CONDITIONS = {
    "attention_ru": (
        "Обращаем ваше внимание, что данные цены на продукцию указаны "
        "за весь предоставленный ассортимент, исполнение данного коммерческого "
        "предложения по частям НЕВОЗМОЖНО."
    ),
    "attention_en": (
        "Please note that these prices for products are indicated for the "
        "entire range provided, and the execution of this commercial offer "
        "is NOT POSSIBLE in parts."
    ),
    "delivery_time_ru": "Срок поставки: 100 дней.",
    "delivery_time_en": "Delivery time: 100 days.",
    "delivery_terms_ru": "Условия доставки: DAP Москва",
    "delivery_terms_en": "Delivery terms: DAP Moscow",
    "payment_terms_ru": "Условия оплаты: в течение 30 дней после выставления счета",
    "payment_terms_en": "Payment terms: 30 days after invoice.",
}


def parse_specification(file_path: str) -> Tuple[Optional[str], Optional[datetime], List[dict]]:
    """
    Читает спецификацию, возвращает (spec_number, spec_date, items).

    ВАЖНО: берётся дата ИЗ САМОЙ СПЕЦИФИКАЦИИ (строка "СПЕЦИФИКАЦИЯ | № 110 от 23.09.2026"),
    а не из строки инвойса. Строка инвойса игнорируется.
    """
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb.active

    spec_number: Optional[str] = None
    spec_date: Optional[datetime] = None

    # Ищем только строку спецификации
    for row in range(1, 12):
        a_val = ws.cell(row, 1).value
        c_val = ws.cell(row, 3).value
        if not a_val or not c_val:
            continue
        a_str = str(a_val).strip()
        c_str = str(c_val).strip()

        # "СПЕЦИФИКАЦИЯ | № 110 от 23.09.2026"
        if "СПЕЦИФИКАЦИЯ" in a_str or "SPECIFICATION" in a_str:
            m = re.search(r'№\s*([^\s]+)\s+от\s+(\d{2}\.\d{2}\.\d{4})', c_str)
            if m:
                spec_number = m.group(1).strip()
                try:
                    spec_date = datetime.strptime(m.group(2), "%d.%m.%Y")
                except ValueError:
                    pass
                break   # ← нашли — больше не ищем

    # Заголовок таблицы
    header_row = None
    for row in range(1, min(ws.max_row, 50) + 1):
        a_val = ws.cell(row, 1).value
        b_val = ws.cell(row, 2).value
        if a_val and str(a_val).strip() == "№":
            if b_val and ("Описание" in str(b_val) or "Description" in str(b_val)):
                header_row = row
                break
    if not header_row:
        raise Exception("Не найден заголовок таблицы в спецификации")

    # Строка Total
    total_row = None
    for row in range(header_row + 1, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            v = ws.cell(row, col).value
            if v and "Total AED" in str(v):
                total_row = row
                break
        if total_row:
            break

    end_row = total_row if total_row else ws.max_row + 1

    # Позиции
    items: List[dict] = []
    for row in range(header_row + 1, end_row):
        num_val = ws.cell(row, 1).value
        desc_val = ws.cell(row, 2).value
        qty_val = ws.cell(row, 8).value
        price_val = ws.cell(row, 9).value

        if not desc_val or not num_val:
            continue

        try:
            qty = int(qty_val) if qty_val else 0
        except (ValueError, TypeError):
            qty = 0
        try:
            price = float(price_val) if price_val else 0.0
        except (ValueError, TypeError):
            price = 0.0

        if qty <= 0:
            continue

        items.append({
            "num": len(items) + 1,
            "description": str(desc_val).strip(),
            "qty": qty,
            "price": price,
        })

    return spec_number, spec_date, items


def generate_commercial_offer(spec_path: str, output_path: str) -> str:
    """
    Генерирует КП по спецификации. Все даты берутся из спецификации.
    """
    spec_number, spec_date, items = parse_specification(spec_path)
    if spec_date is None:
        spec_date = datetime.now()

    # Дата со слэшами — для шапки
    date_str = spec_date.strftime("%d/%m/%Y")
    # Дата с точками — для вводной строки
    date_str_dots = spec_date.strftime("%d.%m.%Y")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "КП"

    # ============ СТИЛИ ============
    font_normal = Font(name="Calibri", size=10)
    font_bold = Font(name="Calibri", size=10, bold=True)
    font_title = Font(name="Calibri", size=12, bold=True)
    border_thin = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin'),
    )
    align_left = Alignment(horizontal='left', vertical='top', wrap_text=True)
    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_right = Alignment(horizontal='right', vertical='center', wrap_text=True)
    header_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")

    # ============ ШАПКА ============
    # Строка 1 — Date / Дата
    ws.cell(1, 1).value = "Date:"
    ws.cell(1, 1).font = font_bold
    ws.cell(1, 2).value = date_str
    ws.cell(1, 2).font = font_bold

    ws.cell(1, 4).value = "Дата:"
    ws.cell(1, 4).font = font_bold
    ws.cell(1, 5).value = date_str
    ws.cell(1, 5).font = font_bold

    # From (RU — колонки A-B, EN — D-E)
    ws.cell(3, 1).value = "От кого:"
    ws.cell(3, 1).font = font_bold
    ws.cell(3, 4).value = "From:"
    ws.cell(3, 4).font = font_bold

    ws.cell(3, 2).value = SUPPLIER["name_ru"]
    ws.cell(3, 2).font = font_normal
    ws.cell(3, 5).value = SUPPLIER["name_en"]
    ws.cell(3, 5).font = font_normal

    ws.cell(4, 2).value = SUPPLIER["address_ru"]
    ws.cell(4, 2).font = font_normal
    ws.cell(4, 5).value = SUPPLIER["address_en"]
    ws.cell(4, 5).font = font_normal

    ws.cell(5, 2).value = SUPPLIER["phone_ru"]
    ws.cell(5, 2).font = font_normal
    ws.cell(5, 5).value = SUPPLIER["phone_en"]
    ws.cell(5, 5).font = font_normal

    # To
    ws.cell(7, 1).value = "Кому:"
    ws.cell(7, 1).font = font_bold
    ws.cell(7, 4).value = "To:"
    ws.cell(7, 4).font = font_bold

    ws.cell(7, 2).value = BUYER["name_ru"]
    ws.cell(7, 2).font = font_normal
    ws.cell(7, 5).value = BUYER["name_en"]
    ws.cell(7, 5).font = font_normal

    ws.cell(8, 2).value = BUYER["address_ru_1"]
    ws.cell(8, 2).font = font_normal
    ws.cell(8, 5).value = BUYER["address_en_1"]
    ws.cell(8, 5).font = font_normal

    ws.cell(9, 2).value = BUYER["address_ru_2"]
    ws.cell(9, 2).font = font_normal
    ws.cell(9, 5).value = BUYER["address_en_2"]
    ws.cell(9, 5).font = font_normal

    ws.cell(10, 2).value = BUYER["phone_ru"]
    ws.cell(10, 2).font = font_normal
    ws.cell(10, 5).value = BUYER["phone_en"]
    ws.cell(10, 5).font = font_normal

    # Заголовок КП
    ws.cell(12, 1).value = "Коммерческое предложение/ Commercial Proposal"
    ws.cell(12, 1).font = font_title
    ws.merge_cells(start_row=12, start_column=1, end_row=12, end_column=5)

    # Вводная строка — даты из спецификации
    ws.cell(14, 1).value = (
        f"В ответ на Ваш запрос № {date_str_dots} от {date_str_dots} "
        f"KVN Group FZCO готово предоставить данный перечень следующих товаров: / "
        f"In response to your request No {date_str_dots} dated {date_str_dots} "
        f"KVN Group FZCO is ready to provide this list of the following goods:"
    )
    ws.cell(14, 1).alignment = align_left
    ws.merge_cells(start_row=14, start_column=1, end_row=14, end_column=5)
    ws.row_dimensions[14].height = 40

    # ============ ТАБЛИЦА ============
    table_header_row = 16
    headers = [
        "№",
        "Описание товара / Description of goods",
        "Ед. / Unit",
        "Количество / QTY",
        "Цена за единицу Дирхам ОАЭ / Unit price AED",
    ]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(table_header_row, col)
        cell.value = h
        cell.font = font_bold
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = border_thin

    # Данные (динамически — сколько позиций, столько строк)
    current_row = table_header_row + 1
    for item in items:
        ws.cell(current_row, 1).value = item["num"]
        ws.cell(current_row, 1).alignment = align_center
        ws.cell(current_row, 1).border = border_thin
        ws.cell(current_row, 1).font = font_normal

        ws.cell(current_row, 2).value = item["description"]
        ws.cell(current_row, 2).alignment = align_left
        ws.cell(current_row, 2).border = border_thin
        ws.cell(current_row, 2).font = font_normal

        ws.cell(current_row, 3).value = "шт/ pcs"
        ws.cell(current_row, 3).alignment = align_center
        ws.cell(current_row, 3).border = border_thin
        ws.cell(current_row, 3).font = font_normal

        ws.cell(current_row, 4).value = item["qty"]
        ws.cell(current_row, 4).alignment = align_center
        ws.cell(current_row, 4).border = border_thin
        ws.cell(current_row, 4).font = font_normal
        ws.cell(current_row, 4).number_format = '#,##0'

        ws.cell(current_row, 5).value = item["price"]
        ws.cell(current_row, 5).alignment = align_right
        ws.cell(current_row, 5).border = border_thin
        ws.cell(current_row, 5).font = font_normal
        ws.cell(current_row, 5).number_format = '#,##0.00'

        current_row += 1

    # Пустая строка после таблицы
    current_row += 1

    # ============ УСЛОВИЯ ============
    ws.cell(current_row, 1).value = CONDITIONS["attention_ru"]
    ws.cell(current_row, 1).alignment = align_left
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)
    ws.row_dimensions[current_row].height = 30
    current_row += 1

    ws.cell(current_row, 1).value = CONDITIONS["attention_en"]
    ws.cell(current_row, 1).alignment = align_left
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)
    ws.row_dimensions[current_row].height = 30
    current_row += 2

    ws.cell(current_row, 1).value = CONDITIONS["delivery_time_ru"]
    current_row += 1
    ws.cell(current_row, 1).value = CONDITIONS["delivery_terms_ru"]
    current_row += 1
    ws.cell(current_row, 1).value = CONDITIONS["payment_terms_ru"]
    current_row += 2

    ws.cell(current_row, 1).value = CONDITIONS["delivery_time_en"]
    current_row += 1
    ws.cell(current_row, 1).value = CONDITIONS["delivery_terms_en"]
    current_row += 1
    ws.cell(current_row, 1).value = CONDITIONS["payment_terms_en"]

    # ============ ШИРИНА КОЛОНОК ============
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 70
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['D'].width = 18
    ws.column_dimensions['E'].width = 22

    # A4, книжная
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4

    wb.save(output_path)
    return output_path